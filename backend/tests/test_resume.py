"""Exercise real document bytes and the public extraction API without model downloads."""

from io import BytesIO
from zipfile import ZIP_DEFLATED, ZipFile

import pytest
from docx import Document
from fastapi.testclient import TestClient
from pypdf import PdfWriter
from pypdf.generic import DecodedStreamObject, DictionaryObject, NameObject

from app import main
from app.resume import (
    MAX_DOCX_UNCOMPRESSED_BYTES,
    MAX_FILE_BYTES,
    MAX_REQUEST_BYTES,
    extract_profile,
)


@pytest.fixture
def client(monkeypatch):
    def no_model():
        raise AssertionError("Resume extraction must not load the embedding model")

    monkeypatch.setattr(main, "get_model", no_model)
    with TestClient(main.app) as test_client:
        yield test_client


def docx_bytes(paragraphs=(), table=None):
    document = Document()
    for text in paragraphs:
        document.add_paragraph(text)
    if table:
        word_table = document.add_table(rows=len(table), cols=len(table[0]))
        for row, values in zip(word_table.rows, table):
            for cell, text in zip(row.cells, values):
                cell.text = text
    target = BytesIO()
    document.save(target)
    return target.getvalue()


def pdf_bytes(lines=(), pages=1, password=None):
    writer = PdfWriter()
    for _ in range(pages):
        page = writer.add_blank_page(width=612, height=792)
        font = DictionaryObject({
            NameObject("/Type"): NameObject("/Font"),
            NameObject("/Subtype"): NameObject("/Type1"),
            NameObject("/BaseFont"): NameObject("/Helvetica"),
        })
        page[NameObject("/Resources")] = DictionaryObject({
            NameObject("/Font"): DictionaryObject({NameObject("/F1"): font}),
        })
        escaped_lines = [line.replace("\\", "\\\\").replace("(", "\\(").replace(")", "\\)") for line in lines]
        commands = "BT /F1 12 Tf 50 740 Td 18 TL\n" + "\n".join(f"({line}) Tj T*" for line in escaped_lines) + "\nET"
        stream = DecodedStreamObject()
        stream.set_data(commands.encode("latin-1"))
        page[NameObject("/Contents")] = writer._add_object(stream)
    if password:
        writer.encrypt(password)
    target = BytesIO()
    writer.write(target)
    return target.getvalue()


RESUME_LINES = [
    "Sam Example",
    "sam@example.com",
    "Education",
    "BS Computer Science, Example University, 2025",
    "Experience",
    "Built a reporting application using Python and PostgreSQL.",
    "Skills",
    "Python, SQL, JavaScript, React, C++, C#, R, Power BI",
    "Interests: Sustainability and accessible technology",
    "Career Goals: Become a data analyst",
]


@pytest.mark.parametrize("extension,make_document,method", [
    ("pdf", pdf_bytes, "pdf_text"),
    ("docx", docx_bytes, "docx_text"),
])
def test_extracts_real_resume_and_stable_contract(client, extension, make_document, method):
    response = client.post("/api/resume/extract", files={"file": (f"resume.{extension}", make_document(RESUME_LINES))})
    assert response.status_code == 200, response.text
    result = response.json()
    assert set(result) == {"filename", "profile", "extracted_text", "warnings", "extraction_method"}
    assert result["extraction_method"] == method
    assert "Computer Science" in result["profile"]["education"]
    assert "reporting application" in result["profile"]["experience"]
    assert {"Python", "SQL", "JavaScript", "React", "C++", "C#", "R", "Power BI", "PostgreSQL"} <= set(result["profile"]["skills"])
    assert result["profile"]["interests"] == "Sustainability and accessible technology"
    assert result["profile"]["career_goals"] == "Become a data analyst"
    assert result["profile"]["work_preferences"] == ""
    assert result["profile"]["preferred_industries"] == ""
    assert "Review and correct" in " ".join(result["warnings"])
    assert "sam@example.com" in result["extracted_text"]


def test_docx_reads_tables_and_preserves_section_boundaries(client):
    data = docx_bytes(["Avery Example"], [
        ["Education", "BSc Physics, 2024"],
        ["Projects", "Built a laboratory reporting tool."],
        ["Technical Skills", "Python, SQL, Microsoft Excel"],
    ])
    result = client.post("/api/resume/extract", files={"file": ("resume.docx", data)}).json()
    assert result["profile"]["education"] == "BSc Physics, 2024"
    assert result["profile"]["experience"] == "Built a laboratory reporting tool."
    assert {"Python", "SQL", "Microsoft Excel"} <= set(result["profile"]["skills"])
    assert "tables" in " ".join(result["warnings"])


def test_skill_matching_uses_whole_tokens_and_does_not_infer_aspirations():
    profile, _ = extract_profile(
        "Sam Example\nhttps://github.com/sam\n"
        "Experience\nManaged events and reacted calmly to problems.\n"
        "No experience with Python.\nI want to learn Docker.\n"
        "Skills\nJavaScript, React, scikit-learn, C++, C#\n"
        "Interests\nPython and Go\nCareer Goals\nLearn Rust and SQL\n"
    )
    assert {"JavaScript", "React", "scikit-learn", "C++", "C#"} <= set(profile.skills)
    assert not {"Java", "C", "R", "Git", "GitHub", "Python", "Docker", "Go", "Rust", "SQL"}.intersection(profile.skills)


def test_taxonomy_skills_are_recognized_and_short_skills_need_explicit_items():
    profile, _ = extract_profile("Skills\nActive Listening; Reading Comprehension; C; R; Go\nExperience\nReady to go wherever needed.")
    assert {"Active Listening", "Reading Comprehension", "C", "R", "Go"} <= set(profile.skills)
    profile, _ = extract_profile("Skills\nReady to go wherever needed.\nExperience\nReceived grade C in R and D.")
    assert not {"C", "R", "Go"}.intersection(profile.skills)


def test_missing_sections_are_empty_and_prompt_manual_review():
    profile, warnings = extract_profile("Sam Example\nA motivated applicant looking for a first opportunity.")
    assert profile.education == profile.experience == ""
    assert profile.skills == []
    assert "education" in " ".join(warnings)
    assert "manually" in " ".join(warnings)


def test_newline_skill_items_preserve_explicit_short_aliases():
    profile, _ = extract_profile("Skills\nPython\nReact\nC\nR\nGo\nExcel")
    assert {"Python", "React", "C", "R", "Go", "Microsoft Excel"} <= set(profile.skills)
    profile, _ = extract_profile("Skills\n- React\n- R\n- Go\n- Excel.")
    assert {"React", "R", "Go", "Microsoft Excel"} <= set(profile.skills)


def test_machine_learning_does_not_negate_following_skills():
    profile, _ = extract_profile("Skills\nMachine learning, Python, SQL, Docker")
    assert {"Machine learning", "Python", "SQL", "Docker"} <= set(profile.skills)
    profile, _ = extract_profile("Experience\nBuilt a machine learning pipeline using Python and SQL.")
    assert {"Machine learning", "Python", "SQL"} <= set(profile.skills)
    profile, _ = extract_profile("Experience\nCurrently learning Python and SQL.")
    assert not {"Python", "SQL"}.intersection(profile.skills)


@pytest.mark.parametrize("filename,data,status,message", [
    ("resume.txt", b"A textual resume", 415, "PDF or DOCX"),
    ("resume.doc", b"A legacy document", 415, "PDF or DOCX"),
    ("resume.pdf", b"", 422, "empty"),
    ("resume.pdf", b"Not really a PDF", 422, "not a readable PDF"),
    ("resume.pdf", b"%PDF-1.4\nbroken content", 422, "could not be read"),
    ("resume.docx", b"Not a ZIP document", 422, "could not be read"),
    ("resume.pdf", b"x" * (MAX_FILE_BYTES + 1), 413, "5 MiB"),
])
def test_upload_validation(client, filename, data, status, message):
    response = client.post("/api/resume/extract", files={"file": (filename, data)})
    assert response.status_code == status, response.text
    assert message in response.json()["detail"]


def test_encrypted_pdf_has_actionable_error(client):
    response = client.post("/api/resume/extract", files={"file": ("resume.pdf", pdf_bytes(RESUME_LINES, password="test-only"))})
    assert response.status_code == 422
    assert "password protected" in response.json()["detail"]


def test_image_only_pdf_has_actionable_error(client):
    response = client.post("/api/resume/extract", files={"file": ("resume.pdf", pdf_bytes())})
    assert response.status_code == 422
    assert "OCR" in response.json()["detail"]


def test_pdf_page_limit(client):
    response = client.post("/api/resume/extract", files={"file": ("resume.pdf", pdf_bytes(RESUME_LINES, pages=31))})
    assert response.status_code == 413
    assert "30 pages" in response.json()["detail"]


def test_docx_uncompressed_size_limit(client):
    target = BytesIO()
    with ZipFile(target, "w", compression=ZIP_DEFLATED) as archive:
        archive.writestr("[Content_Types].xml", "content types")
        archive.writestr("word/document.xml", "x" * (MAX_DOCX_UNCOMPRESSED_BYTES + 1))
    response = client.post("/api/resume/extract", files={"file": ("resume.docx", target.getvalue())})
    assert response.status_code == 413
    assert "expands" in response.json()["detail"]


def test_arbitrary_zip_is_not_accepted_as_docx(client):
    target = BytesIO()
    with ZipFile(target, "w") as archive:
        archive.writestr("example.txt", "This is not a Word document")
    response = client.post("/api/resume/extract", files={"file": ("resume.docx", target.getvalue())})
    assert response.status_code == 422
    assert "not a valid DOCX" in response.json()["detail"]


def test_request_body_limit_before_multipart_parsing(client):
    response = client.post("/api/resume/extract", content=b"x" * (MAX_REQUEST_BYTES + 1))
    assert response.status_code == 413


def test_chunked_request_body_limit(client):
    boundary = "resume-test-boundary"

    def body():
        yield f'--{boundary}\r\nContent-Disposition: form-data; name="file"; filename="resume.pdf"\r\nContent-Type: application/pdf\r\n\r\n'.encode()
        yield b"x" * (MAX_REQUEST_BYTES + 1)
        yield f"\r\n--{boundary}--\r\n".encode()

    response = client.post("/api/resume/extract", content=body(), headers={"Content-Type": f"multipart/form-data; boundary={boundary}"})
    assert response.status_code == 413, response.text


def test_filename_is_sanitized_and_file_field_is_required(client):
    response = client.post("/api/resume/extract", files={"file": ("../../resume.docx", docx_bytes(RESUME_LINES))})
    assert response.status_code == 200
    assert response.json()["filename"] == "resume.docx"
    assert client.post("/api/resume/extract").status_code == 422
